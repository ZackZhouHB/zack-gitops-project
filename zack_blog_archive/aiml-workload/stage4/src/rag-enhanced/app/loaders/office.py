"""Office Document Loaders - Word and Excel"""
import io
import logging
from typing import List
from docx import Document as DocxDocument
from openpyxl import load_workbook

from .base import LoadedDocument

logger = logging.getLogger(__name__)

class WordLoader:
    """Load Microsoft Word (.docx) documents"""
    
    def load(self, file_content: bytes, filename: str) -> LoadedDocument:
        """
        Extract text from Word document.
        Preserves paragraph structure.
        """
        doc = DocxDocument(io.BytesIO(file_content))
        
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # Check if it's a heading
                if para.style and para.style.name.startswith('Heading'):
                    level = para.style.name.replace('Heading ', '')
                    try:
                        hashes = '#' * int(level)
                        paragraphs.append(f"{hashes} {text}")
                    except:
                        paragraphs.append(f"## {text}")
                else:
                    paragraphs.append(text)
        
        # Also extract tables
        for table in doc.tables:
            table_text = self._table_to_markdown(table)
            if table_text:
                paragraphs.append(table_text)
        
        content = "\n\n".join(paragraphs)
        
        logger.info(f"Loaded Word: {filename}, {len(paragraphs)} paragraphs")
        
        return LoadedDocument(
            content=content,
            filename=filename,
            source_type="docx",
            size_bytes=len(file_content),
            checksum=LoadedDocument.calculate_checksum(file_content),
            metadata={"paragraphs": len(paragraphs)}
        )
    
    def _table_to_markdown(self, table) -> str:
        """Convert Word table to markdown"""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
        
        if not rows:
            return ""
        
        lines = []
        # Header
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
        # Data
        for row in rows[1:]:
            lines.append("| " + " | ".join(row) + " |")
        
        return "\n".join(lines)


class ExcelLoader:
    """Load Microsoft Excel (.xlsx) documents - Row-level chunking for better retrieval"""
    
    def load(self, file_content: bytes, filename: str) -> LoadedDocument:
        """
        Extract data from Excel workbook.
        Each row becomes a searchable text block with header context.
        """
        wb = load_workbook(io.BytesIO(file_content), data_only=True)
        
        sections = []
        total_rows = 0
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            
            # Get data as list of rows
            rows = []
            for row in sheet.iter_rows(values_only=True):
                # Skip empty rows
                if any(cell is not None for cell in row):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
            
            if rows and len(rows) > 1:
                header = rows[0]
                total_rows += len(rows) - 1
                
                # Create row-level chunks with header context
                for i, row in enumerate(rows[1:], 1):
                    row_text = f"Sheet: {sheet_name}, Row {i}\n"
                    for col_idx, cell in enumerate(row):
                        if col_idx < len(header) and cell:
                            row_text += f"{header[col_idx]}: {cell}\n"
                    sections.append(row_text.strip())
            elif rows:
                total_rows += len(rows)
                # Convert to markdown table
                section = f"## Sheet: {sheet_name}\n\n"
                section += self._rows_to_markdown(rows)
                sections.append(section)
        
        content = "\n\n".join(sections)
        
        logger.info(f"Loaded Excel: {filename}, {len(wb.sheetnames)} sheets, {total_rows} rows")
        
        return LoadedDocument(
            content=content,
            filename=filename,
            source_type="xlsx",
            size_bytes=len(file_content),
            checksum=LoadedDocument.calculate_checksum(file_content),
            metadata={"sheets": len(wb.sheetnames), "total_rows": total_rows}
        )
    
    def _rows_to_markdown(self, rows: List[List[str]]) -> str:
        """Convert rows to markdown table"""
        if not rows:
            return ""
        
        lines = []
        # Header (first row)
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
        # Data
        for row in rows[1:]:
            # Pad row if needed
            while len(row) < len(rows[0]):
                row.append("")
            lines.append("| " + " | ".join(row[:len(rows[0])]) + " |")
        
        return "\n".join(lines)
